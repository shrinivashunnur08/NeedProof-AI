"""Shop's own sales history from a sheet (POS export, bill book typed into Excel): date, product, qty, price,
optional bill number and customer. Customer names or phone numbers are hashed on upload and never stored.

Real (non-sample) shops get analytics only from this data and from requests in the app; nothing is generated.
"""
from __future__ import annotations

import hashlib
import io
import re
from datetime import date, datetime, timedelta

from app.commerce_loop import store
from app.commerce_loop.inventory import _to_int, _to_price, read_sheet

MAX_SALES_ROWS = 20000
SALES_HEADERS = {
    "date": {"date", "day", "sale date", "bill date", "invoice date", "data", "日期", "तारीख"},
    "product": {"product", "item", "item name", "product name", "description", "produto", "商品", "सामान"},
    "qty": {"qty", "quantity", "units", "count", "quantidade", "数量", "मात्रा"},
    "unit_price": {"price", "unit price", "rate", "mrp", "selling price", "preço", "价格", "दाम"},
    "bill_id": {"bill", "bill no", "bill number", "invoice", "invoice no", "order", "order id", "receipt", "nota"},
    "customer": {"customer", "customer name", "phone", "mobile", "customer phone", "cliente", "顾客", "ग्राहक"},
}


def _headers(headers: list[str]) -> dict[str, str]:
    found: dict[str, str] = {}
    norm = {h: re.sub(r"[^a-z0-9\u0080-￿ ]+", " ", str(h).casefold()).strip() for h in headers}
    for field, aliases in SALES_HEADERS.items():
        for h, n in norm.items():
            if h not in found.values() and (n in aliases or any(n.startswith(a + " ") for a in aliases)):
                found[field] = h
                break
    return found


def _to_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%m/%d/%Y", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    return None


def _anon(shop_id: str, customer) -> str | None:
    s = re.sub(r"\s+", " ", str(customer or "")).strip().casefold()
    return hashlib.sha256(f"{shop_id}|{s}".encode()).hexdigest()[:12] if s else None


def _read(data: bytes, filename: str) -> list[dict]:
    import app.commerce_loop.inventory as inv

    old = inv.MAX_ROWS
    inv.MAX_ROWS = MAX_SALES_ROWS
    try:
        return read_sheet(data, filename)
    finally:
        inv.MAX_ROWS = old


def import_sales(shop_id: str, data: bytes, filename: str) -> dict:
    rows = _read(data, filename)
    if not rows:
        return {"imported": 0, "errors": ["The sheet is empty"], "columns": {}}
    cols = _headers(list(rows[0].keys()))
    missing = [f for f in ("date", "product", "qty") if f not in cols]
    if missing:
        return {"imported": 0, "errors": [f"Missing column: {', '.join(missing)} (see the template)"], "columns": cols}
    shop = store.shop(shop_id) or {}
    currency = None
    from app.commerce_loop.locale import CITIES
    if shop.get("city") in CITIES:
        currency = CITIES[shop["city"]]["currency"]
    good, errors, today = [], [], date.today()
    for n, r in enumerate(rows, start=2):
        d, product, qty = _to_date(r.get(cols["date"])), str(r.get(cols["product"]) or "").strip(), _to_int(r.get(cols["qty"]))
        if d is None or not product or not qty or qty <= 0:
            if len(errors) < 5:
                errors.append(f"Row {n}: needs a date, a product and a quantity above 0")
            continue
        if d > today or d < today - timedelta(days=365):
            if len(errors) < 5:
                errors.append(f"Row {n}: date {d} is in the future or older than a year")
            continue
        price = _to_price(r.get(cols["unit_price"])) if "unit_price" in cols else None
        bill = str(r.get(cols["bill_id"]) or "").strip()[:40] if "bill_id" in cols else ""
        cust = _anon(shop_id, r.get(cols["customer"])) if "customer" in cols else None
        good.append((shop_id, None, None, product[:120], qty, price, currency, f"{d.isoformat()}T12:00:00+00:00",
                     bill or None, cust, "upload"))
    with store._LOCK:  # noqa: SLF001 - one transaction for the whole sheet
        c = store.conn()
        c.execute("DELETE FROM sales WHERE shop_id=? AND source='upload'", (shop_id,))
        c.executemany("INSERT INTO sales (shop_id, inventory_id, request_id, product, qty, unit_price, currency, created_at,"
                      " bill_id, customer_ref, source) VALUES (?,?,?,?,?,?,?,?,?,?,?)", good)
        c.commit()
    days = sorted({g[7][:10] for g in good})
    return {"imported": len(good), "skipped": len(rows) - len(good), "errors": errors, "columns": cols,
            "first_day": days[0] if days else None, "last_day": days[-1] if days else None,
            "customers": len({g[9] for g in good if g[9]}), "bills": len({g[8] for g in good if g[8]}),
            "note": "Customer names/phones are replaced by an anonymous code; a new upload replaces the previous one."}


def template_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Date", "Bill no", "Product", "Qty", "Price", "Customer phone (optional)"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(["2026-09-01", "B-1001", "Your product name", 2, 450, "98xxxxxx01"])
    ws.append(["2026-09-01", "B-1001", "Another product", 1, 120, "98xxxxxx01"])
    ws.append(["2026-09-02", "B-1002", "Your product name", 1, 450, ""])
    notes = wb.create_sheet("How to fill")
    for line in ("One row per product line on a bill. Date, Product and Qty are required.",
                 "Same Bill no on several rows = bought together (used for basket analysis).",
                 "Customer phone or name is optional; it is turned into an anonymous code and never stored.",
                 "Forecast needs 42+ days of sales; customer groups need 24+ customers; basket rules need 24+ bills with 2+ items.",
                 "Replace the example rows with your own sales."):
        notes.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
